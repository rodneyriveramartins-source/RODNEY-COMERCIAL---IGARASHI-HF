# -*- coding: utf-8 -*-
"""Gera DRE_BoxRecifePE_1S2026.xlsx a partir de dre_dados.json.

Abas: RESUMO + CONSOLIDADO + 8 produtos.
Meses = dados da fonte; ACUMULADO e % = fórmulas (recalculáveis).
"""
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

d = json.load(open("dre_dados.json", encoding="utf-8"))
secoes = d["secoes"]

# Paleta Igarashi (hub)
VERDE_ESCURO = "051F20"
TEAL = "0A2E2F"
SALVIA = "8EB69B"
MENTA = "DAF1DE"
DOURADO = "D9C988"

MESES = ["JANEIRO/26", "FEVEREIRO/26", "MARÇO/26", "ABRIL/26", "MAIO/26", "JUNHO/26"]

FMT_BRL = 'R$ #,##0.00;[Red]-R$ #,##0.00'
FMT_PCT = '0.0%;[Red]-0.0%'

f_titulo = Font(name="Arial", size=13, bold=True, color="FFFFFF")
f_sub = Font(name="Arial", size=9, italic=True, color=MENTA)
f_head = Font(name="Arial", size=9, bold=True, color="FFFFFF")
f_normal = Font(name="Arial", size=9)
f_bold = Font(name="Arial", size=9, bold=True)
f_link = Font(name="Arial", size=9, color="008000")  # verde: link entre abas
f_link_b = Font(name="Arial", size=9, bold=True, color="008000")

fill_titulo = PatternFill("solid", fgColor=VERDE_ESCURO)
fill_head = PatternFill("solid", fgColor=TEAL)
fill_total = PatternFill("solid", fgColor=MENTA)
fill_zebra = PatternFill("solid", fgColor="F4FAF5")
fill_gold = PatternFill("solid", fgColor="F5EED5")

thin = Side(style="thin", color="CCDDCC")
borda = Border(left=thin, right=thin, top=thin, bottom=thin)

TOTAIS = ("RECEITA LÍQUIDA", "(-) CMV", "(=) LUCRO BRUTO", "(-) DESPESAS OPERACIONAIS",
          "(=) RESULTADO OPERACIONAL", "(+/-) RESULTADO FINANCEIRO",
          "(=) RESULTADO ANTES DO IRPJ e CSLL", "(=) RESULTADO LÍQUIDO")

wb = Workbook()

def nome_aba(nome):
    return "CONSOLIDADO" if nome == "Consolidado" else nome

def monta_aba(ws, sec):
    nome = nome_aba(sec["nome"])
    linhas = sec["linhas"]
    n = len(linhas)

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"DRE GERENCIAL — {nome} — Box Ceasa Recife/PE — 1º Semestre 2026"
    c.font = f_titulo
    c.fill = fill_titulo
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:H2")
    c = ws["A2"]
    c.value = 'Fonte: "RESULTADO 1S2026 Box RecifePE" (DRE Gerencial, relatório interno Igarashi HF). ACUMULADO e % por fórmula — diferenças de até R$ 0,02 vs. o relatório-fonte são arredondamento dos meses na fonte.'
    c.font = f_sub
    c.fill = fill_titulo
    ws.row_dimensions[2].height = 14

    # cabeçalho da tabela de valores
    hdr = ["LINHA"] + MESES + ["ACUMULADO 1S26"]
    for j, h in enumerate(hdr, start=1):
        cc = ws.cell(row=3, column=j, value=h)
        cc.font = f_head
        cc.fill = fill_head
        cc.alignment = Alignment(horizontal="center" if j > 1 else "left", vertical="center", wrap_text=True)
        cc.border = borda
    ws.row_dimensions[3].height = 22

    receita_row = None
    r0 = 4
    for i, ln in enumerate(linhas):
        r = r0 + i
        eh_total = ln["label"] in TOTAIS
        if ln["label"] == "RECEITA LÍQUIDA":
            receita_row = r
        cc = ws.cell(row=r, column=1, value=ln["label"])
        cc.font = f_bold if eh_total else f_normal
        cc.border = borda
        if eh_total:
            cc.fill = fill_gold if "RESULTADO LÍQUIDO" in ln["label"] else fill_total
        elif i % 2 == 1:
            cc.fill = fill_zebra
        for m in range(6):
            cc = ws.cell(row=r, column=2 + m, value=ln["valores"][m])
            cc.number_format = FMT_BRL
            cc.font = f_bold if eh_total else f_normal
            cc.border = borda
            if eh_total:
                cc.fill = fill_gold if "RESULTADO LÍQUIDO" in ln["label"] else fill_total
            elif i % 2 == 1:
                cc.fill = fill_zebra
        cc = ws.cell(row=r, column=8, value=f"=SUM(B{r}:G{r})")
        cc.number_format = FMT_BRL
        cc.font = f_bold if eh_total else f_normal
        cc.border = borda
        if eh_total:
            cc.fill = fill_gold if "RESULTADO LÍQUIDO" in ln["label"] else fill_total
        elif i % 2 == 1:
            cc.fill = fill_zebra

    # bloco AV% (análise vertical: % da Receita Líquida)
    rp = r0 + n + 1
    ws.merge_cells(start_row=rp, start_column=1, end_row=rp, end_column=8)
    c = ws.cell(row=rp, column=1, value="ANÁLISE VERTICAL — % da Receita Líquida (fórmulas)")
    c.font = f_head
    c.fill = fill_head
    for j, h in enumerate(hdr, start=1):
        cc = ws.cell(row=rp + 1, column=j, value=h)
        cc.font = f_head
        cc.fill = fill_head
        cc.alignment = Alignment(horizontal="center" if j > 1 else "left", vertical="center", wrap_text=True)
        cc.border = borda
    for i, ln in enumerate(linhas):
        r = rp + 2 + i
        rv = r0 + i
        eh_total = ln["label"] in TOTAIS
        cc = ws.cell(row=r, column=1, value=ln["label"])
        cc.font = f_bold if eh_total else f_normal
        cc.border = borda
        if eh_total:
            cc.fill = fill_total
        elif i % 2 == 1:
            cc.fill = fill_zebra
        for m in range(7):
            col = 2 + m
            L = get_column_letter(col)
            cc = ws.cell(row=r, column=col,
                         value=f"=IF({L}${receita_row}=0,0,{L}{rv}/{L}${receita_row})")
            cc.number_format = FMT_PCT
            cc.font = f_bold if eh_total else f_normal
            cc.border = borda
            if eh_total:
                cc.fill = fill_total
            elif i % 2 == 1:
                cc.fill = fill_zebra

    ws.column_dimensions["A"].width = 38
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 16
    ws.freeze_panes = "B4"
    return receita_row

# ---- abas de dados ----
refs = {}  # nome_aba -> (receita_row, n_linhas, labels->row)
ws0 = wb.active
ws0.title = "RESUMO"
for sec in secoes:
    ws = wb.create_sheet(nome_aba(sec["nome"]))
    rrow = monta_aba(ws, sec)
    lab2row = {ln["label"]: 4 + i for i, ln in enumerate(sec["linhas"])}
    refs[nome_aba(sec["nome"])] = (rrow, len(sec["linhas"]), lab2row)

# ---- aba RESUMO ----
ws = ws0
ws.merge_cells("A1:I1")
c = ws["A1"]
c.value = "RESUMO EXECUTIVO — DRE Box Ceasa Recife/PE — 1º Semestre 2026 — Igarashi HF"
c.font = f_titulo
c.fill = fill_titulo
c.alignment = Alignment(vertical="center")
ws.row_dimensions[1].height = 26
ws.merge_cells("A2:I2")
c = ws["A2"]
c.value = 'Todos os valores por fórmula, ligados às abas de detalhe. Fonte: "RESULTADO 1S2026 Box RecifePE" (DRE Gerencial, relatório interno).'
c.font = f_sub
c.fill = fill_titulo

LINHAS_RESUMO = ["RECEITA LÍQUIDA", "(-) CMV", "(=) LUCRO BRUTO",
                 "(-) DESPESAS OPERACIONAIS", "(=) RESULTADO OPERACIONAL",
                 "(+/-) RESULTADO FINANCEIRO", "(=) RESULTADO LÍQUIDO"]

# bloco 1: consolidado mês a mês
r = 4
ws.cell(row=r, column=1, value="CONSOLIDADO — MÊS A MÊS").font = f_head
ws.cell(row=r, column=1).fill = fill_head
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
r += 1
hdr = ["LINHA"] + MESES + ["ACUMULADO 1S26", "% Receita"]
for j, h in enumerate(hdr, start=1):
    cc = ws.cell(row=r, column=j, value=h)
    cc.font = f_head
    cc.fill = fill_head
    cc.border = borda
    cc.alignment = Alignment(horizontal="center" if j > 1 else "left", wrap_text=True)
rrow_c, _, lab2row_c = refs["CONSOLIDADO"]
r += 1
res_liq_row = None
for i, lab in enumerate(LINHAS_RESUMO):
    src = lab2row_c[lab]
    eh_res = lab == "(=) RESULTADO LÍQUIDO"
    if eh_res:
        res_liq_row = r
    cc = ws.cell(row=r, column=1, value=lab)
    cc.font = f_bold
    cc.border = borda
    cc.fill = fill_gold if eh_res else (fill_zebra if i % 2 else PatternFill())
    for m in range(7):
        col = 2 + m
        L = get_column_letter(col)
        cc = ws.cell(row=r, column=col, value=f"=CONSOLIDADO!{L}{src}")
        cc.number_format = FMT_BRL
        cc.font = f_link_b
        cc.border = borda
        if eh_res:
            cc.fill = fill_gold
        elif i % 2:
            cc.fill = fill_zebra
    cc = ws.cell(row=r, column=9, value=f"=IF(H{r - i + 0}=0,0,H{r}/H${r - i})" if i else "=IF(H{0}=0,0,1)".format(r))
    # % receita: divide pelo acumulado da receita (primeira linha do bloco)
    cc.value = f"=IF(H${r - i}=0,0,H{r}/H${r - i})"
    cc.number_format = FMT_PCT
    cc.font = f_bold
    cc.border = borda
    if eh_res:
        cc.fill = fill_gold
    r += 1

# bloco 2: acumulado por produto
r += 1
ws.cell(row=r, column=1, value="ACUMULADO 1S26 — POR PRODUTO (fórmulas ligadas às abas)").font = f_head
ws.cell(row=r, column=1).fill = fill_head
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
r += 1
hdr2 = ["PRODUTO", "RECEITA LÍQUIDA", "CMV", "LUCRO BRUTO", "DESPESAS OPER.",
        "RES. OPERACIONAL", "RES. FINANCEIRO", "RESULTADO LÍQUIDO", "MARGEM LÍQ. %"]
for j, h in enumerate(hdr2, start=1):
    cc = ws.cell(row=r, column=j, value=h)
    cc.font = f_head
    cc.fill = fill_head
    cc.border = borda
    cc.alignment = Alignment(horizontal="center" if j > 1 else "left", wrap_text=True)
r += 1
prod_ini = r
PRODS = [nome_aba(s["nome"]) for s in secoes[1:]]
COLS_MAP = ["RECEITA LÍQUIDA", "(-) CMV", "(=) LUCRO BRUTO", "(-) DESPESAS OPERACIONAIS",
            "(=) RESULTADO OPERACIONAL", "(+/-) RESULTADO FINANCEIRO", "(=) RESULTADO LÍQUIDO"]
for i, p in enumerate(PRODS):
    _, _, l2r = refs[p]
    cc = ws.cell(row=r, column=1, value=p)
    cc.font = f_bold
    cc.border = borda
    if i % 2:
        cc.fill = fill_zebra
    for j, lab in enumerate(COLS_MAP):
        src = l2r.get(lab)
        col = 2 + j
        sheet_ref = f"'{p}'" if " " in p else p
        cc = ws.cell(row=r, column=col, value=f"={sheet_ref}!H{src}")
        cc.number_format = FMT_BRL
        cc.font = f_link
        cc.border = borda
        if i % 2:
            cc.fill = fill_zebra
    cc = ws.cell(row=r, column=9, value=f"=IF(B{r}=0,0,H{r}/B{r})")
    cc.number_format = FMT_PCT
    cc.font = f_bold
    cc.border = borda
    if i % 2:
        cc.fill = fill_zebra
    r += 1
# linha TOTAL (confere com consolidado)
cc = ws.cell(row=r, column=1, value="TOTAL (soma produtos)")
cc.font = f_bold
cc.fill = fill_total
cc.border = borda
for j in range(7):
    col = 2 + j
    L = get_column_letter(col)
    cc = ws.cell(row=r, column=col, value=f"=SUM({L}{prod_ini}:{L}{r - 1})")
    cc.number_format = FMT_BRL
    cc.font = f_bold
    cc.fill = fill_total
    cc.border = borda
cc = ws.cell(row=r, column=9, value=f"=IF(B{r}=0,0,H{r}/B{r})")
cc.number_format = FMT_PCT
cc.font = f_bold
cc.fill = fill_total
cc.border = borda
r += 1
cc = ws.cell(row=r, column=1,
             value="Conferência: a linha TOTAL deve bater com a aba CONSOLIDADO (diferenças de centavos = arredondamento da fonte).")
cc.font = f_sub
cc.font = Font(name="Arial", size=8, italic=True, color="666666")

ws.column_dimensions["A"].width = 30
for col in "BCDEFGHI":
    ws.column_dimensions[col].width = 17

wb.save("DRE_BoxRecifePE_1S2026.xlsx")
print("gerado DRE_BoxRecifePE_1S2026.xlsx")
